"""Fill Naomi's NPO report template (NPO verslag SJABLOON.docx) from a content JSON.

Usage:
    python tools/build_report.py <sjabloon.docx> <content.json> <out.docx> [--html out.html]

The template keeps its letterhead, styles and footer; only text is replaced.
Text between [square brackets] is highlighted yellow: those are the points she
still has to check or complete.

content.json keys (all Dutch text):
    personalia: {naam, geboortedatum, verwijzer, consult, opleiding, leeftijd,
                 beroep, lateraliteit}
    probleemstelling: str
    anamnese, observaties, bespreking: [paragraph, ...]
    besluit_voor: [paragraph, ...]   paragraphs before the standard caveat
    besluit_na:   [paragraph, ...]   diagnosis / advice after the caveat
    onderzoeksdatum: "dd/mm/jjjj"
    footer: {naam, geboortedatum}
    normen: str                      first row of the results table
    rows: [{find, set: {cell_index: text}}]     fill table rows
    delete_rows: [find, ...]                    tests that were not administered
    remarks_after: [{after, text}]              extra *Opmerking rows
A row is found when `find` is a substring of its cell texts joined by " | ".
"""

from __future__ import annotations

import copy
import html
import json
import re
import sys

import docx
from docx.enum.text import WD_COLOR_INDEX

BRACKETS = re.compile(r"(\[[^\]]*\])")


# ---------------------------------------------------------------- helpers

def set_runs(paragraph, text, italic=None, bold=None):
    """Replace paragraph text, keeping the formatting of its first run."""
    runs = paragraph.runs
    proto = copy.deepcopy(runs[0]._r) if runs else None
    for r in runs:
        r._r.getparent().remove(r._r)
    for part in BRACKETS.split(text):
        if not part:
            continue
        if proto is not None:
            new = copy.deepcopy(proto)
            paragraph._p.append(new)
            run = paragraph.runs[-1]
            run.text = part
        else:
            run = paragraph.add_run(part)
        if italic is not None:
            run.italic = italic
        if bold is not None:
            run.bold = bold
        if part.startswith("["):
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def insert_paragraph_after(paragraph, text):
    new_p = copy.deepcopy(paragraph._p)
    paragraph._p.addnext(new_p)
    para = docx.text.paragraph.Paragraph(new_p, paragraph._parent)
    set_runs(para, text)
    return para


def fill_paragraphs(first, texts):
    set_runs(first, texts[0])
    last = first
    for t in texts[1:]:
        last = insert_paragraph_after(last, t)
    return last


def unique_cells(row):
    seen, cells = set(), []
    for c in row.cells:
        if id(c._tc) not in seen:
            seen.add(id(c._tc))
            cells.append(c)
    return cells


def row_key(row):
    return " | ".join(c.text.strip() for c in unique_cells(row))


def find_row(table, needle):
    for row in table.rows:
        if needle in row_key(row):
            return row
    raise KeyError(f"Row not found in template: {needle!r}")


def set_cell(cell, text):
    paras = cell.paragraphs
    set_runs(paras[0], text)
    for p in paras[1:]:
        p._p.getparent().remove(p._p)


def find_paragraph(doc, startswith):
    for p in doc.paragraphs:
        if p.text.strip().startswith(startswith):
            return p
    raise KeyError(f"Paragraph not found in template: {startswith!r}")


# ---------------------------------------------------------------- main fill

def build(template, content, out):
    doc = docx.Document(template)
    c = content
    pers = c["personalia"]

    lines = {
        "Naam:": f"Naam: {pers['naam']}\t\t\t\t\t\tGeboortedatum: {pers['geboortedatum']}",
        "Verwijzer:": f"Verwijzer: {pers['verwijzer']}\t\t\t\t\tConsult: {pers['consult']}",
        "Opleiding:": f"Opleiding: {pers['opleiding']}\t\tLeeftijd: {pers['leeftijd']}",
        "Beroep:": f"Beroep: {pers['beroep']}\t\t\t\t\t{pers['lateraliteit']}",
    }
    for start, text in lines.items():
        set_runs(find_paragraph(doc, start), text)

    set_runs(find_paragraph(doc, "Meneer werd verwezen"), c["probleemstelling"])
    fill_paragraphs(find_paragraph(doc, "Meneer _"), c["anamnese"])
    fill_paragraphs(find_paragraph(doc, "Meneer consulteert"), c["observaties"])

    # Bespreking: the template has six domain paragraphs; replace them all.
    first = find_paragraph(doc, "Psychometrisch onderzoek toont")
    template_paras = []
    p = first
    while not p.text.strip().startswith("Besluit"):
        template_paras.append(p)
        p = docx.text.paragraph.Paragraph(p._p.getnext(), p._parent)
    for extra in template_paras[1:]:
        extra._p.getparent().remove(extra._p)
    fill_paragraphs(first, c["bespreking"])

    # Besluit
    intro = find_paragraph(doc, "Meneer _, een")
    for start in ("Dit neuropsychologisch onderzoek toont", "Observationeel maakt meneer"):
        q = find_paragraph(doc, start)
        q._p.getparent().remove(q._p)
    fill_paragraphs(intro, c["besluit_voor"])
    caveat = find_paragraph(doc, "Bij de interpretatie van deze testresultaten")
    last = caveat
    for t in c["besluit_na"]:
        last = insert_paragraph_after(last, t)

    heading = find_paragraph(doc, "Neuropsychologisch onderzoek (")
    set_runs(heading, f"Neuropsychologisch onderzoek ({c['onderzoeksdatum']})")

    # Tables: results (0) and questionnaires (1)
    tables = doc.tables
    first_cell = unique_cells(tables[0].rows[0])[0]
    set_runs(first_cell.paragraphs[0], "Normen", bold=True)
    set_runs(first_cell.paragraphs[1], c["normen"])

    anchors = {}
    for spec in c.get("remarks_after", []):
        table = next(t for t in tables if any(spec["after"] in row_key(r) for r in t.rows))
        anchors[id(spec)] = (table, find_row(table, spec["after"]))
    # Remove rows of tests that were not administered before matching the rest.
    for needle in c.get("delete_rows", []):
        table = next(t for t in tables if any(needle in row_key(r) for r in t.rows))
        row = find_row(table, needle)
        row._tr.getparent().remove(row._tr)

    # Copy a remark row before any row text is replaced.
    remark_proto = copy.deepcopy(find_row(tables[0], "*Opmerking: Meneer geeft aan")._tr)

    for spec in c["rows"]:
        table = next(t for t in tables if any(spec["find"] in row_key(r) for r in t.rows))
        row = find_row(table, spec["find"])
        cells = unique_cells(row)
        for idx, text in spec["set"].items():
            set_cell(cells[int(idx)], text)

    for spec in c.get("remarks_after", []):
        table, row = anchors[id(spec)]
        new_tr = copy.deepcopy(remark_proto)
        row._tr.addnext(new_tr)
        new_row = docx.table._Row(new_tr, table)
        set_cell(unique_cells(new_row)[0], spec["text"])

    # Footer
    for section in doc.sections:
        for para in section.footer.paragraphs:
            if para.text.startswith("Klinisch neuropsychologisch onderzoeksverslag"):
                f = c["footer"]
                set_runs(para, f"Klinisch neuropsychologisch onderzoeksverslag {f['naam']} "
                               f"(°{f['geboortedatum']})")

    doc.save(out)
    return doc


# ---------------------------------------------------------------- HTML copy

def to_html(doc_path):
    """Simple HTML rendering of the filled report (for a Google Docs copy)."""
    doc = docx.Document(doc_path)
    body = doc.element.body
    out = []

    def esc(t):
        t = html.escape(t)
        return BRACKETS.sub(r'<span style="background:#ffff00">\1</span>', t)

    for child in body.iterchildren():
        tag = child.tag.split("}")[1]
        if tag == "p":
            p = docx.text.paragraph.Paragraph(child, doc)
            text = p.text.strip()
            if not text:
                continue
            if p.style.name.startswith("Heading"):
                out.append(f"<h2>{esc(text)}</h2>")
            elif text == "Besluit":
                out.append(f"<h3>{esc(text)}</h3>")
            else:
                out.append(f"<p>{esc(p.text)}</p>")
        elif tag == "tbl":
            t = docx.table.Table(child, doc)
            out.append('<table border="1" cellpadding="3" style="border-collapse:collapse;font-size:10pt">')
            for row in t.rows:
                cells = unique_cells(row)
                if len(cells) == 1:
                    out.append(f'<tr><td colspan="6"><i>{esc(cells[0].text)}</i></td></tr>')
                else:
                    out.append("<tr>" + "".join(f"<td>{esc(c.text)}</td>" for c in cells) + "</tr>")
            out.append("</table>")
    return ("<html><head><meta charset='utf-8'></head><body style='font-family:Calibri,Arial'>"
            "<h1>Klinisch neuropsychologisch onderzoeksverslag</h1>" + "\n".join(out) + "</body></html>")


if __name__ == "__main__":
    if len(sys.argv) < 4:
        print(__doc__)
        sys.exit(2)
    template, content_path, out_path = sys.argv[1:4]
    content = json.load(open(content_path, encoding="utf-8"))
    build(template, content, out_path)
    if "--html" in sys.argv:
        html_path = sys.argv[sys.argv.index("--html") + 1]
        open(html_path, "w", encoding="utf-8").write(to_html(out_path))
    print(f"Report written to {out_path}")
